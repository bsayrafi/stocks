import pandas as pd
from openpyxl.styles import PatternFill, Font, Alignment
from openpyxl.utils import get_column_letter

def export_df_with_row_colors(
    df, 
    file_path, 
    target_col, 
    color_map, 
    sheet_name="Sheet1",
    header_bg="1F4E78",     # Dark Blue
    header_text="FFFFFF"    # White text
):
    """
    Exports a DataFrame to Excel:
    - Converts values that look like numbers into actual numbers (int/float).
    - Preserves text, percentages, and None values untouched.
    - Styles the header row.
    - Colors data rows based on keyword matches in target_col.
    - Auto-fits column widths.
    """
    fills = {
        str(val).strip().lower(): PatternFill(start_color=hex_code, fill_type="solid")
        for val, hex_code in color_map.items()
    }
    
    with pd.ExcelWriter(file_path, engine="openpyxl") as writer:
        df.to_excel(writer, sheet_name=sheet_name, index=False,na_rep='NaN')
        worksheet = writer.sheets[sheet_name]

        # 1. Convert anything that looks like a number to a number
        non_numeric_tokens = {"nan", "none", "null", "inf", "-inf", "infinity", "-infinity", "n/a", "na", ""}
        for row in range(2, len(df) + 2):
            for col in range(1, len(df.columns) + 1):
                cell = worksheet.cell(row=row, column=col)
                if cell.value is not None:
                    cell_str = str(cell.value).strip()
                    if cell_str.lower() in non_numeric_tokens:
                        continue
                    try:
                        if '.' in cell_str:
                            cell.value = float(cell_str)
                        else:
                            cell.value = int(cell_str)
                    except ValueError:
                        # Keeps text, percentages, and non-numeric strings as-is
                        pass

        # 2. Style Header Row
        if header_bg:
            header_fill = PatternFill(start_color=header_bg, fill_type="solid")
            header_font = Font(name="Calibri", size=11, bold=True, color=header_text)
            for cell in worksheet[1]:
                cell.fill = header_fill
                cell.font = header_font
                cell.alignment = Alignment(horizontal="center", vertical="center")
        
        # 3. Style Data Rows based on target_col keywords
        col_idx = df.columns.get_loc(target_col) + 1
        
        for row in range(2, len(df) + 2):
            cell_text = str(worksheet.cell(row=row, column=col_idx).value or "").strip().lower()
            
            matched_fill = None
            for keyword, fill in fills.items():
                if keyword in cell_text:
                    matched_fill = fill
                    break
            
            if matched_fill:
                for cell in worksheet[row]:
                    cell.fill = matched_fill

        # 4. Auto-size Columns
        for col in worksheet.columns:
            max_len = max(len(str(cell.value or '')) for cell in col)
            col_letter = get_column_letter(col[0].column)
            worksheet.column_dimensions[col_letter].width = max(max_len + 4, 12)